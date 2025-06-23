from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Name in blue
name = doc.add_paragraph()
name_run = name.add_run("UTKARSHA KUNDIKAR")
name_run.font.size = Pt(22)
name_run.font.bold = True
name_run.font.color.rgb = RGBColor(0, 102, 204)  # Blue color
name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Title, size 14 and bold
title = doc.add_paragraph()
title_run = title.add_run("Quality Engineer | Manual Testing | Software Engineer")
title_run.font.size = Pt(14)
title_run.font.bold = True
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Contact Info
contact = doc.add_paragraph()
contact_run = contact.add_run("📧 utkarsha.kundikar2000@gmail.com | 📍 Pune / Mumbai, India | 📞 +91 9373931842")
contact_run.font.size = Pt(10)
contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph()  # Spacer

def add_section_header(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text.upper())
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 102, 204)  # Blue color
    para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT


def add_body_paragraph(doc, text, font_size=10.5):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

def add_bullet_list(doc, items, font_size=10.5):
    for item in items:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(item)
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0, 0, 0)

# Profile Section
add_section_header(doc, "Profile")
add_body_paragraph(doc, "Detail-oriented Quality Engineer with 3 years of experience in web application testing, frontend validation, and UI/UX QA. Proficient in writing and executing test cases, identifying usability issues, and collaborating with cross-functional teams to ensure software quality. Adept in using tools like Cypress (basic level), Jira, and Chrome DevTools for tracking, automation exploration, and debugging. Passionate about learning and open to evolving into a Software Development Engineer in Test (SDET) role.")

# Technical Skills Section
add_section_header(doc, "Technical Skills")
skills = [
    "Languages: JavaScript (ES6+), TypeScript, Python, HTML5, CSS3",
    "Frontend: React.js, Redux, Bootstrap, Tailwind, jQuery",
    "Testing: Cypress (UI tests), Jest (optional)",
    "Backend (Basic understanding): Django, Node.js",
    "Database: MongoDB, PostgreSQL",
    "Dev Tools: Git, Jira, VS Code, Chrome DevTools",
    "Others: RESTful APIs, AI/ML-enhanced workflows, CI/CD (basic), JSON, CSV",
    "Testing: Manual testing (Web UI, usability), Test Plan & Test Case creation, Defect reporting"
]
add_bullet_list(doc, skills)

# Work Experience Section
add_section_header(doc, "Work Experience")
add_body_paragraph(doc, "Jio Platforms Limited – Software Developer / QA-Focused Frontend")
add_body_paragraph(doc, "📍 Navi Mumbai | 🗓️ May 2022 – Present")
jio_experience = [
    "Collaborated with development teams to test and validate user interfaces across multiple web applications.",
    "Conducted extensive manual testing on UI components, user flows, and APIs to identify functional bugs and design mismatches.",
    "Created structured test cases for frontend modules and reported issues using Jira.",
    "Verified responsiveness and cross-browser compatibility for enterprise platforms.",
    "Worked closely with designers and developers to ensure pixel-perfect UI and smooth UX.",
    "Gained hands-on exposure to Cypress for automating basic test scenarios as part of process improvement."
]
add_bullet_list(doc, jio_experience)

add_body_paragraph(doc, "")
add_body_paragraph(doc, "Automate Engineering – IoT Intern")
add_body_paragraph(doc, "📍 Pune | 🗓️ Feb 2022 – Jul 2022")
add_body_paragraph(doc, "Tech Stack: Python, C++, Raspberry Pi, Arduino")
iot_experience = [
    "Developed real-time systems for road safety, including pothole detection and vehicle speed monitoring.",
    "Built a robotic arm controller and a face mask detection system using OpenCV and sensor modules."
]
add_bullet_list(doc, iot_experience)

# Education Section
add_body_paragraph(doc, "")
add_section_header(doc, "Education")
Edjucational_background = [
    "Bachelor of Technology (B.Tech) in Electronics & Telecommunication Engineering",
    "Dr. Babasaheb Ambedkar Technological University – Lonere, Raigad | 📅 2022 | CGPA: 8.6"
]
add_bullet_list(doc, Edjucational_background)

# Achievements Section
add_section_header(doc, "Achievements")
achievements = [
    "Directed a proficient team of 5 members during a competitive hackathon, showcasing leadership and collaborative skills.",
    "Spearheaded cultural initiatives, facilitating diverse events to promote cultural awareness and engagement within the college community."
]
add_bullet_list(doc, achievements)

doc.save("utkarsha_kundikar_resume_4.docx")
